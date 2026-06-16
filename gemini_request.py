from __future__ import annotations
import argparse
import asyncio
import json
import os
import re
import sys
from pathlib import Path
from pptx import Presentation

DEFAULT_MODEL = os.getenv("GEMINI_MODEL", "gemini-2.5-flash")
DEFAULT_CONCURRENCY = max(1, int(os.getenv("GEMINI_CONCURRENCY", "2")))
DEFAULT_MAX_RETRIES = max(1, int(os.getenv("GEMINI_MAX_RETRIES", "4")))
DEFAULT_BACKOFF_SECONDS = float(os.getenv("GEMINI_BACKOFF_SECONDS", "2.0"))

def clean_text(text):
    if not text:
        return ""
    text = text.replace("\u00a0", " ")
    text = re.sub(r"\s+", " ", text)
    return text.strip()

def extract_slide_text(slide):
    parts = []
    for shape in slide.shapes:
        if getattr(shape, "has_text_frame", False):
            raw = getattr(shape, "text", "")
            cleaned = clean_text(raw)
            if cleaned:
                parts.append(cleaned)
    return "\n".join(parts).strip()

def extract_slides_from_pptx(pptx_path):
    prs = Presentation(str(pptx_path))
    slides = []
    for i, slide in enumerate(prs.slides, start=1):
        text = extract_slide_text(slide)
        if not text:
            continue
        slides.append({
            "slide_number": i,
            "text": text
        })
    return slides

def build_prompt(slide_number, slide_text):
    return f"""Please explain the following slide as if teaching a beginner software developer:
"Slide {slide_number}:
{slide_text}"
""".strip()

def is_retryable_error(exc):
    msg = str(exc).lower()
    return any(x in msg for x in ["429", "rate limit", "503", "500", "timeout", "resource_exhausted"])

def extract_retry_delay(err_msg):
    match_seconds = re.search(r"seconds:\s*(\d+)", err_msg)
    if match_seconds:
        return float(match_seconds.group(1)) + 1.5
    match_retry_in = re.search(r"retry in\s*([\d.]+)\s*s", err_msg, re.IGNORECASE)
    if match_retry_in:
        return float(match_retry_in.group(1)) + 1.5
    return 0.0

class MockGeminiExplainer:
    async def explain_slide(self, slide_number, slide_text):
        await asyncio.sleep(0)
        return f"Mock explanation for slide {slide_number}"

class GeminiExplainer:
    def __init__(self, model_name: str = DEFAULT_MODEL):
        import google.generativeai as genai
        api_key = os.getenv("GEMINI_API_KEY")
        if not api_key:
            raise RuntimeError("Missing GEMINI_API_KEY environment variable")
        genai.configure(api_key=api_key)
        self.model = genai.GenerativeModel(model_name)
    async def explain_slide(self, slide_number, slide_text):
        prompt = build_prompt(slide_number, slide_text)
        response = await self.model.generate_content_async(prompt)
        return response.text or ""

async def explain_with_retry(explainer, slide, max_retries, backoff_seconds):
    slide_number = slide["slide_number"]
    slide_text = slide["text"]
    last_error = ""
    for attempt in range(max_retries):
        try:
            explanation = await explainer.explain_slide(slide_number, slide_text)
            return {
                "slide_number": slide_number,
                "text": slide_text,
                "explanation": explanation,
                "error": None
            }
        except Exception as e:
            last_error = str(e)
            if not is_retryable_error(e) or attempt == max_retries - 1:
                return {
                    "slide_number": slide_number,
                    "text": slide_text,
                    "explanation": "",
                    "error": last_error
                }
            dynamic_delay = extract_retry_delay(last_error)
            if dynamic_delay > 0:
                wait_time = dynamic_delay
                print(
                    f"[Rate Limit] Slide {slide_number}: Google requested a pause. Waiting for {wait_time:.1f} seconds...",
                    file=sys.stderr)
            else:
                wait_time = backoff_seconds * (2 ** attempt)
                print(f"[Error] Slide {slide_number}: Temporary error. Retrying in {wait_time:.1f} seconds...",
                      file=sys.stderr)
            await asyncio.sleep(wait_time)

async def process_presentation(pptx_path, explainer):
    slides = extract_slides_from_pptx(pptx_path)
    semaphore = asyncio.Semaphore(DEFAULT_CONCURRENCY)
    async def worker(slide):
        async with semaphore:
            return await explain_with_retry(
                explainer,
                slide,
                DEFAULT_MAX_RETRIES,
                DEFAULT_BACKOFF_SECONDS
            )
    return await asyncio.gather(*(worker(s) for s in slides))

def save_results(results, output_path):
    output_path.write_text(
        json.dumps(results, ensure_ascii=False, indent=2),
        encoding="utf-8"
    )

async def run(pptx_file, use_mock=False):
    pptx_path = Path(pptx_file)
    explainer = MockGeminiExplainer() if use_mock else GeminiExplainer()
    results = await process_presentation(pptx_path, explainer)
    output_path = pptx_path.with_suffix(".json")
    save_results(results, output_path)
    return output_path

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("pptx_path", help="Path to the PowerPoint presentation")
    parser.add_argument("--mock", action="store_true", help="Use Mock explainer instead of calling the API")
    args = parser.parse_args()
    pptx_path = Path(args.pptx_path)
    if not pptx_path.exists():
        print(f"Error: The presentation file '{pptx_path}' was not found.", file=sys.stderr)
        return 1
    if pptx_path.suffix.lower() != ".pptx":
        print(f"Error: Invalid file type ({pptx_path.suffix}). The script only supports .pptx files.", file=sys.stderr)
        return 1
    try:
        output = asyncio.run(run(pptx_path, args.mock))
        print(f"Successfully saved explanations to: {output}")
    except Exception as e:
        print("Error during execution:", e, file=sys.stderr)
        return 1
    return 0

if __name__ == "__main__":
    raise SystemExit(main())