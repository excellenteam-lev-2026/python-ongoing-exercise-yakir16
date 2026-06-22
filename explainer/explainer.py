import os
import re
import sys
from pathlib import Path

BASE_DIR = Path(__file__).resolve().parent.parent
sys.path.append(str(BASE_DIR))

import time
import json
import asyncio
import functools
from datetime import datetime
import requests
from pptx import Presentation
from db.database import SessionLocal, Upload



DEFAULT_MODEL = os.getenv("GEMINI_MODEL", "gemini-3.1-flash-lite")
DEFAULT_CONCURRENCY = max(1, int(os.getenv("GEMINI_CONCURRENCY", "4")))
DEFAULT_MAX_RETRIES = max(1, int(os.getenv("GEMINI_MAX_RETRIES", "4")))
DEFAULT_BACKOFF_SECONDS = float(os.getenv("GEMINI_BACKOFF_SECONDS", "1.0"))
SLEEP_INTERVAL = 10  

UPLOAD_FOLDER = BASE_DIR / "uploads"
OUTPUT_FOLDER = BASE_DIR / "outputs"
UPLOAD_FOLDER.mkdir(exist_ok=True)
OUTPUT_FOLDER.mkdir(exist_ok=True)

def clean_text(text: str):
    if not text:
        return ""
    text = text.replace("\u00a0", " ")
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def extract_slide_text(slide):
    parts = []
    for shape in slide.shapes:
        if getattr(shape, "has_text_frame", False):
            raw = getattr(shape, "text", "")
            cleaned = clean_text(raw)
            if cleaned:
                parts.append(cleaned)
    return "\n".join(parts).strip()


def extract_slides_from_pptx(pptx_path):
    prs = Presentation(str(pptx_path))
    slides = []
    for i, slide in enumerate(prs.slides, start=1):
        text = extract_slide_text(slide)
        if not text:
            continue
        slides.append({"slide_number": i, "text": text})
    return slides


def build_prompt(slide_number,slide_text):
    return f"Explain this slide simply for a beginner software developer:\n\nSlide {slide_number}:\n{slide_text}".strip()


class GeminiExplainer:
    def __init__(self, model_name: str = DEFAULT_MODEL):
        self.api_key = os.getenv("GEMINI_API_KEY")
        if not self.api_key:
            raise RuntimeError("Missing GEMINI_API_KEY")
        self.model_name = model_name
        self.url = f"https://generativelanguage.googleapis.com/v1beta/models/{self.model_name}:generateContent?key={self.api_key}"

    async def explain_slide(self,slide_number,slide_text):
        prompt = build_prompt(slide_number, slide_text)
        payload = {
            "contents": [{
                "parts": [{"text": prompt}]
            }]
        }
        headers = {"Content-Type": "application/json"}

        loop = asyncio.get_running_loop()
        func = functools.partial(requests.post, self.url, json=payload, headers=headers)
        response = await loop.run_in_executor(None, func)
        response.raise_for_status()
        
        response_data = response.json()
        return response_data['candidates'][0]['content']['parts'][0]['text']


async def explain_with_retry(explainer, slide, max_retries, backoff_seconds):
    slide_number = slide["slide_number"]
    slide_text = slide["text"]
    last_error = ""
    for attempt in range(max_retries):
        try:
            explanation = await explainer.explain_slide(slide_number, slide_text)
            return {
                "slide_number": slide_number,
                "text": slide_text,
                "explanation": explanation,
                "error": None
            }
        except Exception as e:
            last_error = str(e)
            if attempt == max_retries - 1:
                return {
                    "slide_number": slide_number,
                    "text": slide_text,
                    "explanation": "",
                    "error": last_error
                }
            
            if hasattr(e, 'response') and e.response is not None and e.response.status_code == 429:
                sleep_time = 30 * (attempt + 1)
                print(f"[{time.strftime('%X')}] Rate limit (429) hit on slide {slide_number}. Sleeping for {sleep_time}s...")
                await asyncio.sleep(sleep_time)
            else:
                await asyncio.sleep(backoff_seconds * (2 ** attempt))


async def process_presentation(pptx_path, explainer):
    slides = extract_slides_from_pptx(pptx_path)
    semaphore = asyncio.Semaphore(DEFAULT_CONCURRENCY)
    async def worker(slide):
        async with semaphore:
            return await explain_with_retry(
                explainer,
                slide,
                DEFAULT_MAX_RETRIES,
                DEFAULT_BACKOFF_SECONDS
            )
    return await asyncio.gather(*(worker(s) for s in slides))


async def main_loop():
    print(f"[{time.strftime('%X')}] Explainer started in DB mode. Polling for pending uploads...")
    explainer = GeminiExplainer()

    while True:
        try:
            with SessionLocal() as db:
                upload = db.query(Upload).filter(Upload.status == "pending").order_by(Upload.upload_time.asc()).first()

                if upload:
                    print(f"[{time.strftime('%X')}] DEBUG START: Processing new file: {upload.filename} (UID: {upload.uid})")
                    upload.status = "processing"
                    db.commit()
                    try:
                        results = await process_presentation(upload.get_upload_path(), explainer)
                        upload.get_output_path().write_text(
                            json.dumps(results, ensure_ascii=False, indent=2),
                            encoding="utf-8"
                        )
                        upload.status = "done"
                        upload.get_upload_path().unlink(missing_ok=True)
                        print(f"[{time.strftime('%X')}] DEBUG END: Finished processing UID: {upload.uid}")
                        
                    except Exception as e:
                        upload.status = "failed"
                        upload.error_message = str(e)
                        print(f"[{time.strftime('%X')}] Error processing UID {upload.uid}: {e}", file=sys.stderr)
                        
                    finally:
                        upload.finish_time = datetime.utcnow()
                        db.commit()
                        
        except Exception as e:
            print(f"Error in main loop: {e}", file=sys.stderr)
            
        await asyncio.sleep(SLEEP_INTERVAL)


if __name__ == "__main__":
    asyncio.run(main_loop())